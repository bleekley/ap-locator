"""Offline regressions: python -m unittest discover -s src -p 'test_*.py'."""
import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

import collect
import config
import correlate
import mac_trace
import parsers
import test_parsers

FIX = Path(__file__).resolve().parent.parent / "fixtures"


class ReadinessTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.addCleanup(patch.stopall)
        patch.object(config, "OUTPUT", self.root).start()

    def inventory(self, body):
        path = self.root / "switches.csv"
        path.write_text("hostname,mgmt_ip,street_address,owner\n" + body)
        return path, correlate.load_inventory(path)

    def test_aireos_existing_fixtures(self):
        test_parsers.main()
        text = (FIX / "aireos_ap_summary.txt").read_text()
        self.assertEqual(len(parsers.validated_summary(text, parsers.parse_aireos_ap_summary)), 3)

    def test_9800_collection_fixture(self):
        conn = Mock()
        conn.send_command.side_effect = [(FIX / "9800_ap_summary_SYNTHETIC.txt").read_text(),
                                         (FIX / "9800_cdp_SYNTHETIC.txt").read_text()]
        with patch.object(collect, "ConnectHandler", return_value=conn):
            rows = collect.collect_9800("lab", "unused", {})
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["neighbor_ip"], "10.0.1.1")
        self.assertEqual(rows[1]["eth_mac"], "aabb.ccdd.0002")
        conn.disconnect.assert_called_once()

    def test_rejected_or_truncated_summary(self):
        good = (FIX / "9800_ap_summary_SYNTHETIC.txt").read_text()
        for text in ["% Authorization failed", "unrecognized output", good.replace("APs: 2", "APs: 3")]:
            with self.subTest(text=text), self.assertRaises(ValueError):
                parsers.validated_summary(text, parsers.parse_9800_ap_summary)
        self.assertEqual(parsers.validated_summary("Number of APs: 0", parsers.parse_9800_ap_summary), [])

    def test_cdp_command_failure_disconnects(self):
        conn = Mock()
        conn.send_command.side_effect = [(FIX / "9800_ap_summary_SYNTHETIC.txt").read_text(), "% Invalid input detected"]
        with patch.object(collect, "ConnectHandler", return_value=conn), self.assertRaises(ValueError):
            collect.collect_9800("lab", "unused", {})
        conn.disconnect.assert_called_once()

    def test_controller_validation(self):
        path = self.root / "controllers.csv"
        bad = ["name,hosts,type\na,1,9800\n", "name,host,type\na,,9800\n",
               "name,host,type\n", "name,host,type\na,1,9800\na,2,aireos\n",
               "name,host,type\na,1,9800\nb,1,9800\n", "name,host,type\na,1,bad\n"]
        for body in bad:
            path.write_text(body)
            with self.subTest(body=body), self.assertRaises(SystemExit):
                config.controllers(path)
        path.write_text("name,host,type\n a , 1 , 9800 \n")
        self.assertEqual(config.controllers(path), [{"name": "a", "host": "1", "type": "9800"}])

    def test_matching_rules_and_extras(self):
        _, inv = self.inventory("sw-a,10.0.1.1,100 Main,Ops\n")
        for ap, method in [({"neighbor_ip": "10.0.1.1"}, "neighbor IP"),
                           ({"neighbor_name": "SW-A"}, "hostname"),
                           ({"neighbor_name": "sw-a.example.org"}, "hostname (domain stripped)")]:
            with self.subTest(ap=ap):
                result = correlate.match_row(pd.Series(ap), inv)
                self.assertEqual(result["location_status"], "matched")
                self.assertEqual(result["match_method"], method)
                self.assertEqual(result["inventory_owner"], "Ops")

    def test_blank_keys_do_not_match(self):
        _, inv = self.inventory("sw-a,,100 Main,Ops\n,10.0.1.2,200 Main,Ops\n")
        for blank in ["", float("nan"), None, pd.NA]:
            result = correlate.match_row(pd.Series({"neighbor_name": blank, "neighbor_ip": blank}), inv)
            self.assertEqual(result["location_status"], "unmatched")

    def test_duplicate_matches_and_missing_address(self):
        for body, ap, expected in [
            ("sw-a,1,100 Main,A\nsw-a,2,200 Main,B\n", {"neighbor_name": "sw-a"}, "ambiguous inventory match"),
            ("sw-a,1,100 Main,A\nsw-b,1,200 Main,B\n", {"neighbor_ip": "1", "neighbor_name": "sw-a"}, "ambiguous inventory match"),
            ("sw-a,1,,A\n", {"neighbor_ip": "1"}, "missing street address"),
        ]:
            _, inv = self.inventory(body)
            result = correlate.match_row(pd.Series(ap), inv)
            self.assertEqual(result["location_status"], expected)
            self.assertEqual(result["street_address"], correlate.UNKNOWN)

    def run_collection(self, results):
        controllers = [{"name": f"wlc{i}", "host": str(i), "type": "9800"} for i in range(len(results))]
        with patch.object(config, "controllers", return_value=controllers), \
             patch.object(config, "creds_for", return_value={}), \
             patch.object(collect, "collect_9800", side_effect=results):
            if any(isinstance(r, Exception) for r in results):
                with self.assertRaises(SystemExit) as e:
                    collect.main()
                self.assertEqual(e.exception.code, 1)
            else:
                collect.main()

    def report(self, path, *options):
        with patch.object(sys, "argv", ["correlate.py", str(path), *options]):
            correlate.main()

    def test_partial_collection_requires_opt_in_and_labels_workbook(self):
        ap = parsers.merge("wlc0", parsers.parse_9800_ap_summary((FIX / "9800_ap_summary_SYNTHETIC.txt").read_text()), {})
        self.run_collection([ap, RuntimeError("offline")])
        path, _ = self.inventory("sw-a,1,100 Main,Ops\n")
        with self.assertRaises(SystemExit):
            self.report(path)
        self.assertFalse((self.root / "ap_location_report.xlsx").exists())
        self.report(path, "--allow-partial")
        book = load_workbook(self.root / "ap_location_report.xlsx")
        self.assertEqual(book["Run Summary"]["A2"].value, "INCOMPLETE - controllers failed")
        self.assertEqual(book["Controller Status"]["D3"].value, "failed")
        summary = pd.read_excel(self.root / "ap_location_report.xlsx", sheet_name="Per-Site Summary")
        self.assertEqual(summary.ap_count.sum(), 2)
        self.assertEqual(summary.street_address.iloc[0], correlate.UNKNOWN)
        book.close()

    def test_missing_address_stays_in_report_totals(self):
        aps = parsers.parse_9800_ap_summary((FIX / "9800_ap_summary_SYNTHETIC.txt").read_text())
        cdp = parsers.parse_9800_cdp((FIX / "9800_cdp_SYNTHETIC.txt").read_text())
        self.run_collection([parsers.merge("wlc0", aps, cdp)])
        path, _ = self.inventory("sw-a,10.0.1.1,,Ops\nsw-b,10.0.1.2,200 Main,Facilities\n")
        self.report(path)
        report = self.root / "ap_location_report.xlsx"
        site = pd.read_excel(report, sheet_name="Per-Site Summary")
        self.assertEqual(site.ap_count.sum(), 2)
        self.assertIn(correlate.UNKNOWN, site.street_address.tolist())
        ap_rows = pd.read_excel(report, sheet_name="AP Locations")
        self.assertEqual(ap_rows.location_status.tolist(), ["missing street address", "matched"])
        self.assertEqual(ap_rows.inventory_owner.tolist(), ["Ops", "Facilities"])

    def test_invalid_config_preserves_previous_inventory(self):
        target = self.root / "ap_inventory.csv"
        target.write_text("previous inventory")
        with patch.object(config, "controllers", side_effect=SystemExit("invalid list")), \
             patch.object(collect, "ConnectHandler") as connect, self.assertRaises(SystemExit):
            collect.main()
        connect.assert_not_called()
        self.assertEqual(target.read_text(), "previous inventory")

    def test_empty_success_report(self):
        self.run_collection([[]])
        path, _ = self.inventory("sw-a,1,100 Main,Ops\n")
        self.report(path)
        book = load_workbook(self.root / "ap_location_report.xlsx")
        self.assertEqual(book["Run Summary"]["D2"].value, 0)
        self.assertEqual(book["AP Locations"].max_row, 1)
        book.close()

    def test_all_failed_report_is_blocked(self):
        self.run_collection([RuntimeError("offline")])
        path, _ = self.inventory("sw-a,1,100 Main,Ops\n")
        with self.assertRaises(SystemExit):
            self.report(path)

    def test_modified_or_legacy_inventory_is_blocked(self):
        self.run_collection([[]])
        path, _ = self.inventory("sw-a,1,100 Main,Ops\n")
        with (self.root / "ap_inventory.csv").open("a") as f:
            f.write("\n")
        with self.assertRaises(SystemExit):
            self.report(path)
        (self.root / "collection_status.json").unlink()
        with self.assertRaises(SystemExit):
            self.report(path)

    def test_mac_ambiguity_and_cli_error(self):
        conn = Mock()
        conn.send_command.return_value = "1 aabb.ccdd.eeff DYNAMIC Gi1/0/1\n2 aabb.ccdd.eeff DYNAMIC Gi1/0/2"
        with self.assertRaises(ValueError):
            mac_trace.find_port(conn, "aabb.ccdd.eeff")
        conn.send_command.return_value = "% Authorization failed"
        with self.assertRaises(ValueError):
            mac_trace.cdp_switch_on_port(conn, "Gi1/0/1")

    def test_multihop_trace_stays_unresolved_without_endpoint_evidence(self):
        core, edge = Mock(), Mock()
        core.find_prompt.return_value = "core#"
        edge.find_prompt.return_value = "edge#"
        core.send_command.side_effect = ["1 aabb.ccdd.eeff DYNAMIC Gi1/0/48", "Device ID: edge\nIP address: 10.0.0.2\nCapabilities: Switch"]
        edge.send_command.side_effect = ["1 aabb.ccdd.eeff DYNAMIC Gi1/0/1", "Total cdp entries displayed : 0"]
        output = io.StringIO()
        with patch.object(mac_trace, "ConnectHandler", side_effect=[core, edge]), \
             patch.object(config, "creds", return_value={}), \
             patch.object(sys, "argv", ["mac_trace.py", "aabb.ccdd.eeff", "10.0.0.1"]), \
             contextlib.redirect_stdout(output), self.assertRaises(SystemExit) as e:
            mac_trace.main()
        self.assertIn("UNRESOLVED", str(e.exception))
        self.assertIn("edge", str(e.exception))
        self.assertNotIn("FOUND", output.getvalue())
        core.disconnect.assert_called_once()
        edge.disconnect.assert_called_once()


if __name__ == "__main__":
    unittest.main()
